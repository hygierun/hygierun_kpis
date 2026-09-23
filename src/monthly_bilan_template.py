"""Génère le fichier Excel 'template' d'Input Mensuel consolidé : une feuille par feuille attendue,
avec les colonnes obligatoires (+ les colonnes de détail utiles) déjà en en-tête - reste à l'utilisateur
d'aller coller les données export ERP dans chaque feuille (et de renommer les colonnes datées comme
'T [date export]' ou '08/26 HT' selon le mois traité).
"""

from io import BytesIO
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .monthly_loader import DETAIL_COLUMNS, bilan_total_required_columns

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

# Stock_Invent : colonnes réelles nécessaires au calcul au-delà des colonnes obligatoires (Dépot,
# Famille) - la colonne "T <date>" est dynamique (retrouvée par son préfixe "T "), à renommer chaque
# mois avec la date du jour de l'export (ex. "T 04/09/26").
EXTRA_COLUMNS: Dict[str, List[str]] = {
    "Stock_Invent": ["Désignation", "Référence", "Ref. fourn.", "Unité", "T [date export]", "Nb jours", "Anc Ref"],
}


def _merge_columns(*groups: List[str]) -> List[str]:
    merged: List[str] = []
    for group in groups:
        for col in group:
            if col not in merged:
                merged.append(col)
    return merged


def generate_bilan_template_excel(year: int, month: int) -> bytes:
    """Fichier Excel vierge (en-têtes seulement) prêt à être rempli pour le mois donné."""
    required = bilan_total_required_columns(year, month)
    wb = Workbook()
    wb.remove(wb.active)

    for sheet, cols in required.items():
        columns = _merge_columns(cols, DETAIL_COLUMNS.get(sheet, []), EXTRA_COLUMNS.get(sheet, []))
        ws = wb.create_sheet(sheet)
        for col_idx, name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=name)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            ws.column_dimensions[cell.column_letter].width = max(14, min(28, len(str(name)) + 4))
        ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
