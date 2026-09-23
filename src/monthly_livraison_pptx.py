"""Remplit le template PowerPoint Livraison + Compta (3 diapos) avec les KPI calculés."""

from pathlib import Path

from pptx import Presentation

from .monthly_commerce import evolution
from .pptx_helpers import (COMMENT_PLACEHOLDER, find_shape, fr, fr_k, set_cell, set_delta, set_delta_na,
                           set_subtitle, set_text)

MULTI_BL_TABLE_ROWS = ("0", "1", "2", "3", "4 et plus")


def _pct(value) -> str:
    return "n/a" if value is None else f"{fr(value, 1 if value < 10 else 0)}%"


def _set_label_and_value(shape, label, value: str):
    """Boîte 'libellé / retour à la ligne / valeur' : on ne touche que le run du libellé (2e) et la valeur (dernier run)."""
    runs = shape.text_frame.paragraphs[0].runs
    if label is not None:
        runs[1].text = label
    runs[-1].text = value


def _fill_slide1(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, n1, m1 = d["cur"], d["n1"], d["m1"]
    cam, cam1, cam_m1 = cur["camions"], n1["camions"], m1["camions"]

    set_text(find_shape(slide, 9), fr_k(cam["ca_livre"]))
    set_delta(find_shape(slide, 11), f"N-1 : {fr_k(cam1['ca_livre'])}", e["ca_livre_n1"])
    set_text(find_shape(slide, 28), f"{cam['part_ca_livre']:.0f}% du CA livré")

    set_text(find_shape(slide, 22), f"{fr(cam['bla_moy'], 1)} BLA")
    set_delta(find_shape(slide, 24), f"N-1 : {fr(cam1['bla_moy'], 1)} BLA", e["bla_n1"])
    set_text(find_shape(slide, 25), f"{fr(cam['fact_moy'], 1)} Fact")
    set_delta(find_shape(slide, 27), f"N-1 : {fr(cam1['fact_moy'], 1)} Fact", e["fact_n1"])

    enl, enl1 = cur["enlevements"], n1["enlevements"]
    set_text(find_shape(slide, 51), str(enl["enlevements"]))
    set_delta(find_shape(slide, 18), f"N-1 : {enl1['enlevements']}", e["enlevements_n1"])
    set_text(find_shape(slide, 75), str(enl["commerciaux"]))
    set_delta(find_shape(slide, 16), f"N-1 : {enl1['commerciaux']}", e["commerciaux_n1"])

    set_text(find_shape(slide, 48), str(cam["microstor"]))
    set_text(find_shape(slide, 54), f"Mois-1 : {cam_m1['microstor']}")
    set_text(find_shape(slide, 87), str(cam["sous_traitance"]))
    set_text(find_shape(slide, 57), f"Mois-1 : {cam_m1['sous_traitance']}")

    # Blocs GPS
    gps, gps_m1 = cur.get("gps"), m1.get("gps")
    table = find_shape(slide, 88).table
    if gps is None:
        for shape_id in (80, 30):
            set_text(find_shape(slide, shape_id), "n/a")
        for shape_id in (36, 44):
            set_delta_na(find_shape(slide, shape_id), "Mois-1")
        for row in range(1, len(table.rows)):
            for col in range(1, len(table.columns)):
                set_cell(table.cell(row, col), "n/a")
    else:
        pond = gps.loc["Moy pond"]
        set_text(find_shape(slide, 80), fr(pond["Nb Arrêts Moy"], 1))
        set_text(find_shape(slide, 30), f"{pond['Distance Moy (kms)']:.0f} kms")
        if gps_m1 is not None:
            pond_m1 = gps_m1.loc["Moy pond"]
            set_delta(find_shape(slide, 36), f"Mois-1 : {fr(pond_m1['Nb Arrêts Moy'], 1)} ", evolution(pond["Nb Arrêts Moy"], pond_m1["Nb Arrêts Moy"]))
            set_delta(find_shape(slide, 44), f"Mois-1 : {pond_m1['Distance Moy (kms)']:.0f} kms ", evolution(pond["Distance Moy (kms)"], pond_m1["Distance Moy (kms)"]))
        else:
            set_delta_na(find_shape(slide, 36), "Mois-1")
            set_delta_na(find_shape(slide, 44), "Mois-1")
        for row_idx, chauffeur in enumerate(gps.index, 1):
            r = gps.loc[chauffeur]
            set_cell(table.cell(row_idx, 0), chauffeur)
            set_cell(table.cell(row_idx, 1), fr(r["Nb Arrêts Moy"], 1))
            set_cell(table.cell(row_idx, 2), f"{r['Distance Moy (kms)']:.0f}")
            set_cell(table.cell(row_idx, 3), str(int(r["Nb de jours travail"])))

    set_text(find_shape(slide, 61), COMMENT_PLACEHOLDER)


def _fill_slide2(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, m1 = d["cur"], d["m1"]

    set_text(find_shape(slide, 9), f"       {fr(cur['delais']['cde_livraison'], 1)} jours")
    set_delta(find_shape(slide, 11), f"Mois-1 : {fr(m1['delais']['cde_livraison'], 1)} jours",
              e["delai_cde_livraison_m1"], higher_is_bad=True)

    multi, multi_m1 = cur["multiples"], m1["multiples"]
    set_text(find_shape(slide, 66), f"         {multi['part_1bl']:.0f} %")
    set_delta(find_shape(slide, 48), f"Mois-1 : {multi_m1['part_1bl']:.0f}%", e["part_1bl_m1"])

    table = find_shape(slide, 54).table
    for row_idx, bucket in enumerate(MULTI_BL_TABLE_ROWS, 1):
        nb = multi["buckets"][bucket]
        set_cell(table.cell(row_idx, 0), bucket)
        set_cell(table.cell(row_idx, 1), str(nb))
        set_cell(table.cell(row_idx, 2), _pct(nb / multi["total"] * 100 if multi["total"] else None))


def _fill_slide3(slide, result: dict):
    year = result["year"]
    d, e = result["data"], result["evolutions"]
    cur, m1 = d["cur"], d["m1"]
    cb, fd = cur["clients_bloques"], cur["factures_dues"]

    m1_cb, m1_fd = m1.get("clients_bloques"), m1.get("factures_dues")

    set_text(find_shape(slide, 94), f"{cb['total']} clients")
    set_text(find_shape(slide, 96), f"Dont {cb['actifs']} actifs en {year}")
    _set_label_and_value(find_shape(slide, 98), f" CA {year}", fr_k(cb["ca"], 1))
    _set_label_and_value(find_shape(slide, 104), None, fr_k(cb["solde"], 1))
    if m1_cb:
        set_delta(find_shape(slide, 107), f"Mois-1 : {m1_cb['total']} clients", e.get("clients_bloques_total_m1"))
    else:
        set_delta_na(find_shape(slide, 107), "Mois-1")

    set_text(find_shape(slide, 72), f"{fd['nb']} fact")
    set_text(find_shape(slide, 74), fr_k(fd["montant"], 1))
    set_text(find_shape(slide, 78), _pct(fd["part_annee_nb"]))
    set_text(find_shape(slide, 80), _pct(fd["part_annee_montant"]))
    if m1_fd:
        set_delta(find_shape(slide, 89), f"Mois-1 : {m1_fd['nb']} fact", e.get("factures_dues_nb_m1"), higher_is_bad=True)
        set_delta(find_shape(slide, 92), f"Mois-1 : {fr_k(m1_fd['montant'], 1)}", e.get("factures_dues_montant_m1"), higher_is_bad=True)
    else:
        set_delta_na(find_shape(slide, 89), "Mois-1")
        set_delta_na(find_shape(slide, 92), "Mois-1")

    set_text(find_shape(slide, 15), "Factures impayées (≥ 60 jrs)")
    set_text(find_shape(slide, 39), f"{fd['nb_60']} fact")
    set_text(find_shape(slide, 28), fr_k(fd["montant_60"], 1))
    set_text(find_shape(slide, 35), f"{fd['clients_60']} clients distincts")
    set_text(find_shape(slide, 82), _pct(fd["part_60_nb"]))
    set_text(find_shape(slide, 84), _pct(fd["part_60_montant"]))
    if m1_fd:
        set_delta(find_shape(slide, 65), f"Mois-1 : {m1_fd['nb_60']} fact", e.get("factures_dues_60_nb_m1"), higher_is_bad=True)
        set_delta(find_shape(slide, 68), f"Mois-1 : {fr_k(m1_fd['montant_60'], 1)}", e.get("factures_dues_60_montant_m1"), higher_is_bad=True)
        if m1_fd.get("clients_60") is not None:
            set_delta(find_shape(slide, 61), f"Mois-1 : {m1_fd['clients_60']} clients", e.get("factures_dues_60_clients_m1"), higher_is_bad=True)
        else:
            set_delta_na(find_shape(slide, 61), "Mois-1")
    else:
        for shape_id in (65, 68, 61):
            set_delta_na(find_shape(slide, shape_id), "Mois-1")

    set_text(find_shape(slide, 24), f"{fr(cur['delais']['livraison_facture'], 1)} jours")
    set_delta(find_shape(slide, 38), f"Mois-1 : {fr(m1['delais']['livraison_facture'], 1)}j",
              e["delai_livraison_facture_m1"], higher_is_bad=True)


def generate_livraison_pptx(result: dict, template_path: str, output_path: str) -> str:
    prs = Presentation(template_path)
    slide1, slide2, slide3 = prs.slides
    for slide in (slide1, slide2, slide3):
        set_subtitle(find_shape(slide, 5), result["year"], result["month"])
    _fill_slide1(slide1, result)
    _fill_slide2(slide2, result)
    _fill_slide3(slide3, result)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path
