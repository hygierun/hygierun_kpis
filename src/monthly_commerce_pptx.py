"""Remplit le template PowerPoint Commerce (3 diapos) avec les KPI calculés."""

from copy import deepcopy
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu

from .monthly_commerce import SEUILS_LIVRAISON, month_bounds
from .monthly_loader import MONTHS_FR

GREEN = RGBColor(46, 125, 50)
RED = RGBColor(192, 57, 43)
GREY = RGBColor(138, 155, 163)
COMMENT_PLACEHOLDER = "Commentaire du mois : à compléter"

# Colonnes du tableau "en attente" : libellé, nb, montant avec Franck, montant hors Franck
EN_ATTENTE_COL_WIDTHS = [1589405, 450000, 1000000, 1000000]
EN_ATTENTE_TABLE_LEFT = 540000


def fr(value: float, decimals: int = 0) -> str:
    return f"{value:.{decimals}f}".replace(".", ",")


def fr_k(value: float, decimals: int = 0) -> str:
    return f"{fr(value / 1000, decimals)} k€"


def fr_pct(value: Optional[float], decimals: int = 0) -> str:
    """'+13 %' / '−28 %' ; les décimales nulles sont supprimées (+12,0 -> +12)."""
    if value is None:
        return "n/a"
    text = fr(abs(value), decimals)
    if decimals and text.endswith("," + "0" * decimals):
        text = text[: -(decimals + 1)]
    sign = "+" if value > 0 else ("−" if value < 0 else "")
    return f"{sign}{text} %"


def arrow(value: Optional[float]) -> str:
    if value is None or value == 0:
        return "►"
    return "▲" if value > 0 else "▼"


def delta_color(value: Optional[float]) -> RGBColor:
    if value is None or value == 0:
        return GREY
    return GREEN if value > 0 else RED


def find_shape(slide, shape_id: int):
    def rec(shapes):
        for shape in shapes:
            if shape.shape_id == shape_id:
                return shape
            if shape.shape_type == 6:
                found = rec(shape.shapes)
                if found is not None:
                    return found
        return None

    shape = rec(slide.shapes)
    if shape is None:
        raise KeyError(f"Forme {shape_id} introuvable dans le template")
    return shape


def _set_paragraph(paragraph, text: str, color: Optional[RGBColor] = None):
    runs = paragraph.runs
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    if color is not None:
        runs[0].font.color.rgb = color


def set_text(shape, text: str):
    paragraphs = shape.text_frame.paragraphs
    _set_paragraph(paragraphs[0], text)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_delta(shape, reference: str, pct: Optional[float], decimals: int = 0):
    """Ligne de référence (ex. 'N-1 : 860 cdes') + ligne d'évolution colorée."""
    paragraphs = shape.text_frame.paragraphs
    delta = f"{arrow(pct)} {fr_pct(pct, decimals)}"
    if len(paragraphs) >= 2:
        _set_paragraph(paragraphs[0], reference)
        _set_paragraph(paragraphs[1], delta, delta_color(pct))
        return
    # Cas d'un seul paragraphe : "référence <saut de ligne> flèche + valeur" en plusieurs runs
    runs = paragraphs[0].runs
    runs[0].text = reference
    runs[1].text, runs[1].font.color.rgb = arrow(pct), delta_color(pct)
    runs[2].text = " "
    runs[3].text, runs[3].font.color.rgb = fr_pct(pct, decimals), delta_color(pct)


def set_subtitle(shape, year: int, month: int):
    runs = shape.text_frame.paragraphs[0].runs
    runs[1].text = MONTHS_FR[month - 1]
    runs[2].text = f" {year}"


def set_cell(cell, text: str):
    paragraph = cell.text_frame.paragraphs[0]
    if paragraph.runs:
        _set_paragraph(paragraph, text)
    else:
        paragraph.add_run().text = text


def _fill_slide1(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, n1 = d["cur"], d["n1"]
    cmd, cmd1 = cur["commandes"], n1["commandes"]
    fac, fac1 = cur["factures"], n1["factures"]

    set_text(find_shape(slide, 9), f"{cmd['nb']} cdes")
    set_delta(find_shape(slide, 11), f"N-1 : {cmd1['nb']} cdes", e["commandes_nb_n1"])
    set_text(find_shape(slide, 66), f"{cmd['nb_arch']} arch. + {cmd['nb_alivr']} à livrer")
    set_text(find_shape(slide, 12), fr_k(cmd["ca"]))
    set_delta(find_shape(slide, 14), f"N-1 : {fr_k(cmd1['ca'])}", e["commandes_ca_n1"])
    set_text(find_shape(slide, 72), f"{fr_k(cmd['ca_arch'])} arch. + {fr_k(cmd['ca_alivr'])} à livrer")

    set_text(find_shape(slide, 110), f"{fac['nb']} fact")
    set_delta(find_shape(slide, 114), f"N-1 : {fac1['nb']}", e["factures_nb_n1"])
    set_text(find_shape(slide, 19), fr_k(fac["ca"]))
    set_delta(find_shape(slide, 21), f"N-1 : {fr_k(fac1['ca'])}", e["factures_ca_n1"])

    set_text(find_shape(slide, 75), f"{fr(fac['pm_avec'], 1)} €")
    set_delta(find_shape(slide, 60), f"N-1 : {fr(fac1['pm_avec'], 1)} €", e["pm_avec_n1"])
    set_text(find_shape(slide, 84), f"{fr(fac['pm_sans'], 1)} €")
    set_delta(find_shape(slide, 88), f"N-1 : {fr(fac1['pm_sans'], 1)} €", e["pm_sans_n1"])

    set_text(find_shape(slide, 22), COMMENT_PLACEHOLDER)


def _add_en_attente_column(slide):
    """Ajoute la colonne 'hors Franck' au tableau des commandes en attente et cale les en-têtes."""
    shape = find_shape(slide, 78)
    tbl = shape._element.graphic.graphicData.tbl
    grid = tbl.tblGrid
    new_col = deepcopy(grid.findall(qn("a:gridCol"))[-1])
    for ext in new_col.findall(qn("a:extLst")):
        new_col.remove(ext)
    grid.append(new_col)
    for tr in tbl.findall(qn("a:tr")):
        last_tc = tr.findall(qn("a:tc"))[-1]
        last_tc.addnext(deepcopy(last_tc))

    for column, width in zip(shape.table.columns, EN_ATTENTE_COL_WIDTHS):
        column.width = Emu(width)
    shape.width = Emu(sum(EN_ATTENTE_COL_WIDTHS))
    shape.left = Emu(EN_ATTENTE_TABLE_LEFT)

    left = shape.left
    lefts = [left + sum(EN_ATTENTE_COL_WIDTHS[:i]) for i in range(4)]
    header_nb, header_net = find_shape(slide, 8), find_shape(slide, 7)
    header_hors = deepcopy(header_net._element)
    header_net._element.addnext(header_hors)
    new_id = max(s.shape_id for s in slide.shapes) + 1
    cnv = header_hors.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
    cnv.set("id", str(new_id))
    cnv.set("name", "ZoneTexte hors Franck")
    header_hors = find_shape(slide, new_id)

    for header, col in ((header_nb, 1), (header_net, 2), (header_hors, 3)):
        header.left, header.width = Emu(lefts[col]), Emu(EN_ATTENTE_COL_WIDTHS[col])
    set_text(header_net, "Avec Franck")
    set_text(header_hors, "Hors Franck")


def _fill_slide2(slide, result: dict):
    year, month = result["year"], result["month"]
    d, e = result["data"], result["evolutions"]
    cur, m1 = d["cur"], d["m1"]
    _, end = month_bounds(year, month)
    month_name = MONTHS_FR[month - 1]

    set_text(find_shape(slide, 32), f"Commandes de {year} en attente de livraison au {end.day} {month_name}  (attente de stock)")
    set_text(find_shape(slide, 69), f"=> Montants = A livrer net / Cdes à livrer / filtre sur date Livraison du 01/01/{year} au {end.strftime('%d/%m/%Y')} / filtre Repr = commerciaux x6")
    set_text(find_shape(slide, 23), f"Livraison > Archivées > filtre {month_name} > filtre Total HT < 100 & 150€")

    _add_en_attente_column(slide)
    table = find_shape(slide, 78).table
    for row_idx, key in enumerate(("initiale", "reliquats", "total"), 0):
        block = cur["en_attente"][key]
        set_cell(table.cell(row_idx, 1), str(block["nb"]))
        set_cell(table.cell(row_idx, 2), fr_k(block["net_avec"], 1))
        set_cell(table.cell(row_idx, 3), fr_k(block["net_sans"], 1))
    set_text(find_shape(slide, 9), COMMENT_PLACEHOLDER)

    set_text(find_shape(slide, 43), f"{cur['nouveaux_clients']['nb']} clients")
    set_text(find_shape(slide, 46), fr_k(cur["nouveaux_clients"]["ca"], 1))
    set_delta(find_shape(slide, 49), f"Mois -1 : {m1['nouveaux_clients']['nb']}", e["nouveaux_clients_nb_m1"], 1)
    set_delta(find_shape(slide, 57), f"Mois -1 : {fr_k(m1['nouveaux_clients']['ca'], 1)}", e["nouveaux_clients_ca_m1"], 1)

    seuil_shapes = {100: {"value": 17, "share": 61, "n1": 21, "m1": 55}, 150: {"value": 34, "share": 63, "n1": 50, "m1": 59}}
    for seuil in SEUILS_LIVRAISON:
        ids = seuil_shapes[seuil]
        s = cur["seuils"]
        set_text(find_shape(slide, ids["value"]), f"         {s[f'nb_{seuil}']}")
        _set_paragraph(find_shape(slide, ids["share"]).text_frame.paragraphs[0], f"{s[f'part_{seuil}']:.0f}% cdes clients")
        for key, label in (("n1", "N-1"), ("m1", "Mois-1")):
            ref = d[key]["seuils"]
            set_delta(find_shape(slide, ids[key]), f"{label} : {ref[f'nb_{seuil}']} = {ref[f'part_{seuil}']:.0f}%",
                      e[f"seuil_{seuil}_{key}"], 1)


def _fill_slide3(slide, result: dict):
    top = result["data"]["cur"]["top_ventes"]
    for table_id, key, value_fmt in ((26, "valeur", lambda r: fr_k(r["HT"], 1)), (10, "volume", lambda r: str(int(round(r["Quantité"]))))):
        table = find_shape(slide, table_id).table
        rows = top[key]
        for i in range(1, len(table.rows)):
            if i - 1 < len(rows):
                row = rows.iloc[i - 1]
                values = [row["Article réf"], row["Article"], value_fmt(row), f"{fr(row['part_ca'], 1)}%"]
            else:
                values = ["", "", "", ""]
            for col, text in enumerate(values):
                set_cell(table.cell(i, col), text)
    set_text(find_shape(slide, 7), COMMENT_PLACEHOLDER)


def generate_commerce_pptx(result: dict, template_path: str, output_path: str) -> str:
    prs = Presentation(template_path)
    slide1, slide2, slide3 = prs.slides
    for slide in (slide1, slide2, slide3):
        set_subtitle(find_shape(slide, 5), result["year"], result["month"])
    _fill_slide1(slide1, result)
    _fill_slide2(slide2, result)
    _fill_slide3(slide3, result)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path
