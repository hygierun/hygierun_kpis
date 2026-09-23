"""Aides communes au remplissage des templates PowerPoint (formats FR, flèches d'évolution, édition de formes)."""

from typing import Optional

from pptx.dml.color import RGBColor

from .monthly_loader import MONTHS_FR

GREEN = RGBColor(46, 125, 50)
RED = RGBColor(192, 57, 43)
GREY = RGBColor(138, 155, 163)
COMMENT_PLACEHOLDER = "Commentaire du mois : à compléter"


def fr(value: Optional[float], decimals: int = 0) -> str:
    """N-1/Mois-1 peuvent être None quand la période de référence n'a aucune donnée (ex. pas de
    tournées ce mois-là l'année dernière) - on l'affiche alors comme les autres évolutions absentes."""
    if value is None:
        return "n/a"
    return f"{value:.{decimals}f}".replace(".", ",")


def fr_k(value: Optional[float], decimals: int = 0) -> str:
    if value is None:
        return "n/a"
    return f"{fr(value / 1000, decimals)} k€"


def fr_pct(value: Optional[float]) -> str:
    """Évolution avec 2 chiffres significatifs : '+5,0 %', '−5,2 %', '−41 %' (0 décimale dès 10 %)."""
    if value is None:
        return "n/a"
    text = fr(abs(value), 1 if abs(value) < 10 else 0)
    sign = "+" if value > 0 else ("−" if value < 0 else "")
    return f"{sign}{text} %"


def arrow(value: Optional[float]) -> str:
    if value is None or value == 0:
        return "►"
    return "▲" if value > 0 else "▼"


def delta_color(value: Optional[float], higher_is_bad: bool = False) -> RGBColor:
    if value is None or value == 0:
        return GREY
    return RED if (value > 0) == higher_is_bad else GREEN


def find_shape(slide, shape_id: int):
    def rec(shapes):
        for shape in shapes:
            if shape.shape_id == shape_id:
                return shape
            if shape.shape_type == 6:
                found = rec(shape.shapes)
                if found is not None:
                    return found
        return None

    shape = rec(slide.shapes)
    if shape is None:
        raise KeyError(f"Forme {shape_id} introuvable dans le template")
    return shape


def set_paragraph(paragraph, text: str, color: Optional[RGBColor] = None):
    runs = paragraph.runs
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    if color is not None:
        runs[0].font.color.rgb = color


def set_text(shape, text: str):
    paragraphs = shape.text_frame.paragraphs
    set_paragraph(paragraphs[0], text)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_delta(shape, reference: str, pct: Optional[float], higher_is_bad: bool = False):
    """Ligne de référence (ex. 'N-1 : 860 cdes') + évolution colorée (higher_is_bad : une hausse est en rouge)."""
    delta = f"{arrow(pct)} {fr_pct(pct)}"
    _write_delta(shape, reference, delta, delta_color(pct, higher_is_bad))


def set_delta_na(shape, label: str):
    """Référence non calculable (photo du jour sans historique) : 'Mois-1 : n/a'."""
    _write_delta(shape, f"{label} : n/a", "—", GREY)


def _write_delta(shape, reference: str, delta: str, color: RGBColor):
    paragraphs = shape.text_frame.paragraphs
    if len(paragraphs) >= 2:
        set_paragraph(paragraphs[0], reference)
        set_paragraph(paragraphs[1], delta, color)
        return
    # Un seul paragraphe (référence, saut de ligne, évolution) : 1er run = référence, 2e run = évolution
    runs = paragraphs[0].runs
    runs[0].text = reference
    runs[1].text, runs[1].font.color.rgb = delta, color
    for extra in runs[2:]:
        extra._r.getparent().remove(extra._r)


def set_subtitle(shape, year: int, month: int):
    runs = shape.text_frame.paragraphs[0].runs
    runs[1].text = MONTHS_FR[month - 1]
    runs[2].text = f" {year}"


def set_cell(cell, text: str):
    paragraph = cell.text_frame.paragraphs[0]
    if paragraph.runs:
        set_paragraph(paragraph, text)
    else:
        paragraph.add_run().text = text
