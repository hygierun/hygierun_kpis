"""Aides communes à la génération des Excel de sortie du Bilan Mensuel."""

from typing import List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, color="366092")

SUMMARY_HEADERS = ["Indicateur", "Valeur", "N-1", "Évol. N-1 (%)", "Mois-1", "Évol. Mois-1 (%)"]


def write_summary(ws, title: str, rows: List[Optional[list]]):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(SUMMARY_HEADERS)
    for col in range(1, len(SUMMARY_HEADERS) + 1):
        cell = ws.cell(row=3, column=col)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL

    for row in rows:
        if row is None:
            ws.append([])
        elif len(row) == 1:
            ws.append(row)
            ws.cell(row=ws.max_row, column=1).font = SECTION_FONT
        else:
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
            for col in (2, 3, 5):
                ws.cell(row=ws.max_row, column=col).number_format = "#,##0.0" if isinstance(row[col - 1], float) else "#,##0"
            for col in (4, 6):
                ws.cell(row=ws.max_row, column=col).number_format = "+0.0;-0.0;0.0"

    ws.column_dimensions["A"].width = 52
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16


def write_frame(wb: Workbook, title: str, df: pd.DataFrame):
    ws = wb.create_sheet(title)
    for col_idx, name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(50, len(str(name)) + 4))
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, value in enumerate(row, 1):
            if isinstance(value, float) and pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if hasattr(value, "year"):
                cell.number_format = "dd/mm/yyyy"


def select_columns(df: pd.DataFrame, wanted: List[str]) -> pd.DataFrame:
    return df[[c for c in wanted if c in df.columns]]
