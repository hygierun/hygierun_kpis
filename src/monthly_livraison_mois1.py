"""Lit les KPI 'photo du jour' (clients bloqués, factures dues, GPS) depuis le PowerPoint Livraison & Compta
généré par cette app le mois précédent, pour servir de référence Mois-1 (ces KPI ne se recalculent pas depuis
l'historique du fichier input). Best-effort : un champ introuvable ou mal formé est simplement omis.
"""

import re
from typing import Optional

from pptx import Presentation

from .pptx_helpers import find_shape


def _text(slide, shape_id: int) -> Optional[str]:
    try:
        return find_shape(slide, shape_id).text_frame.text
    except (KeyError, IndexError, AttributeError):
        return None


def _nombre(text: Optional[str]) -> Optional[float]:
    """Premier nombre du texte (ex. '154 clients' -> 154, '96,6 %' -> 96.6)."""
    if not text:
        return None
    match = re.search(r"[-+]?\d[\d\s]*(?:,\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group().replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _montant_k(text: Optional[str]) -> Optional[float]:
    """Montant en euros à partir d'un texte terminé par 'k€' (ex. 'Dont CA 2026\\x0b236,1 k€' -> 236100).

    Le groupe capturé exclut volontairement \\s : un saut de ligne interne au texte (les cartes 'libellé
    <br> valeur' du template) collerait sinon un nombre voisin (ex. l'année) au montant.
    """
    if not text:
        return None
    match = re.search(r"(\d+(?:,\d+)?)\s*k€", text)
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", ".")) * 1000
    except ValueError:
        return None


def read_mois1_reference(pptx_path: str) -> dict:
    """Retourne un dict partiel {'clients_bloques': {...}, 'factures_dues': {...}, 'gps': {...}}."""
    reference: dict = {}
    try:
        prs = Presentation(pptx_path)
        slide1, _slide2, slide3 = prs.slides
    except Exception:
        return reference

    total = _nombre(_text(slide3, 94))
    ca = _montant_k(_text(slide3, 98))
    if total is not None:
        reference["clients_bloques"] = {"total": int(total), "ca": ca}

    nb = _nombre(_text(slide3, 72))
    montant = _montant_k(_text(slide3, 74))
    nb_60 = _nombre(_text(slide3, 39))
    montant_60 = _montant_k(_text(slide3, 28))
    clients_60 = _nombre(_text(slide3, 35))
    if nb is not None and montant is not None:
        reference["factures_dues"] = {
            "nb": int(nb), "montant": montant,
            "nb_60": int(nb_60) if nb_60 is not None else None,
            "montant_60": montant_60,
            "clients_60": int(clients_60) if clients_60 is not None else None,
        }

    try:
        table = find_shape(slide1, 88).table
        gps = {}
        for row_idx in range(1, len(table.rows)):
            chauffeur = table.cell(row_idx, 0).text.strip()
            arrets, distance, jours = (_nombre(table.cell(row_idx, col).text) for col in (1, 2, 3))
            if chauffeur and arrets is not None:
                gps[chauffeur] = {"Nb Arrêts Moy": arrets, "Distance Moy (kms)": distance, "Nb de jours travail": jours}
        if gps:
            reference["gps"] = gps
    except (KeyError, IndexError, AttributeError):
        pass

    return reference
