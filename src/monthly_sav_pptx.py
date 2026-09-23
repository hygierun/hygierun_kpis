"""Remplit le template PowerPoint SAV/Achat (2 diapos) avec les KPI calculés.

Diapo 1 (SAV) : Devis, Facturation vs Docs Nuls, Main d'œuvre, Déplacement, Productivité et Nombre
d'interventions. Diapo 2 (Achat/Appro) : Valorisation du stock et Articles à épuisement — le reste de la
diapo (Commandes fournisseur, Couverture de stock) n'a pas encore été défini avec Antoine.
"""

from copy import deepcopy
from pathlib import Path

from pptx.util import Pt
from pptx import Presentation

from .monthly_sav import STOCK_FAMILLES
from .pptx_helpers import (GREEN, GREY, RED, arrow, delta_color, find_shape, fr, fr_k, fr_pct, set_cell,
                           set_delta, set_delta_na, set_paragraph, set_subtitle, set_text)

PRODUCTIVITE_SEUIL = 50  # % : vert au-dessus, rouge en dessous (règle validée avec Antoine)


def _heures(value: float) -> str:
    """Pas de décimale pour un nombre d'heures entier (144), une décimale sinon (107,5)."""
    return fr(value, 0) if float(value).is_integer() else fr(value, 1)


def _set_devis_mois1(shape, nb: int, pct):
    """Cas particulier : 3 runs ('Mois', ' -1 : N devis', 'flèche %'), pas le schéma standard à 2 runs."""
    runs = shape.text_frame.paragraphs[0].runs
    runs[1].text = f" -1 : {nb} devis"
    runs[2].text = f"{arrow(pct)} {fr_pct(pct)}"
    runs[2].font.color.rgb = delta_color(pct)


def _set_interventions_total_mois1(shape, value, pct):
    """Cas particulier : 3 runs ('Mois', ' -1 : N', 'flèche %'), comme le Mois-1 des devis (shape 83)."""
    runs = shape.text_frame.paragraphs[0].runs
    if value is None:
        runs[1].text, runs[2].text, runs[2].font.color.rgb = " -1 : n/a", "—", GREY
    else:
        runs[1].text = f" -1 : {value}"
        runs[2].text, runs[2].font.color.rgb = f"{arrow(pct)} {fr_pct(pct)}", delta_color(pct)


def _set_productivite_total_mois1(shape, value, pct):
    """Cas particulier : 2 runs seulement ('Mois', ' -1 :'), il manque le 3e run (flèche %) à créer."""
    paragraph = shape.text_frame.paragraphs[0]
    runs = paragraph.runs
    delta_run = runs[2] if len(runs) > 2 else paragraph.add_run()
    if value is None:
        runs[1].text, delta_run.text, delta_run.font.color.rgb = " -1 : n/a", "—", GREY
    else:
        runs[1].text = f" -1 : {_pct_abs(value)}"
        delta_run.text, delta_run.font.color.rgb = f"{arrow(pct)} {fr_pct(pct)}", delta_color(pct)


def _ensure_second_paragraph(shape, template_shape):
    """La case Mois-1 du CA Docs Nuls n'a qu'un paragraphe (placeholder vide) : on clone la 2e ligne
    (grande, colorée) d'une case N-1/Mois-1 déjà stylée avant d'y écrire l'évolution normalement."""
    if len(shape.text_frame.paragraphs) < 2:
        modele = deepcopy(template_shape.text_frame.paragraphs[1]._p)
        shape.text_frame.paragraphs[0]._p.addnext(modele)


def _fill_devis(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, n1, m1 = d["cur"]["devis"], d["n1"]["devis"], d["m1"]["devis"]

    set_text(find_shape(slide, 100), f"{cur['nb']} devis")
    set_text(find_shape(slide, 102), fr_k(cur["ca"], 1))
    set_delta(find_shape(slide, 93), f"N-1 : {n1['nb']} devis", e["devis_nb_n1"])
    _set_devis_mois1(find_shape(slide, 83), m1["nb"], e["devis_nb_m1"])


def _fill_docs_nuls(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, n1, m1 = d["cur"]["docs_nuls"], d["n1"]["docs_nuls"], d["m1"]["docs_nuls"]

    set_text(find_shape(slide, 112), f"{cur['total']} Fact")
    set_text(find_shape(slide, 118), f"{cur['dn']} DN")
    set_text(find_shape(slide, 120), f"{cur['non_dn']} non DN")
    set_text(find_shape(slide, 104), fr_k(cur["ca"], 1))

    set_delta(find_shape(slide, 129), f"Mois-1 : {m1['total']}", e["docs_nuls_total_m1"])
    set_delta(find_shape(slide, 122), f"Mois-1 : {m1['dn']} DN", e["docs_nuls_dn_m1"])
    set_delta(find_shape(slide, 116), f"Mois-1 : {m1['non_dn']}", e["docs_nuls_non_dn_m1"])

    ca_mois1 = find_shape(slide, 148)
    ca_n1 = find_shape(slide, 106)
    _ensure_second_paragraph(ca_mois1, ca_n1)
    set_delta(ca_n1, f"N-1 : {fr_k(n1['ca'], 1)}", e["docs_nuls_ca_n1"])
    set_delta(ca_mois1, f"Mois-1 : {fr_k(m1['ca'], 1)}", e["docs_nuls_ca_m1"])


def _fill_main_oeuvre(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, n1 = d["cur"]["main_oeuvre"], d["n1"]["main_oeuvre"]

    set_text(find_shape(slide, 72), f"{_heures(cur['total'])}h")
    set_text(find_shape(slide, 78), f"{_heures(cur['ebc'])} EBC")
    set_text(find_shape(slide, 79), f"{_heures(cur['sav'])} SAV")
    set_delta(find_shape(slide, 82), f"N-1 : {_heures(n1['total'])}", e["main_oeuvre_total_n1"])
    set_delta(find_shape(slide, 80), f"N-1 : {_heures(n1['ebc'])}", e["main_oeuvre_ebc_n1"])
    set_delta(find_shape(slide, 77), f"N-1 : {_heures(n1['sav'])}", e["main_oeuvre_sav_n1"])


def _fill_deplacement(slide, result: dict):
    d, e = result["data"], result["evolutions"]
    cur, n1 = d["cur"]["deplacement"], d["n1"]["deplacement"]

    set_text(find_shape(slide, 130), f"{cur['total']:.0f} u")
    set_text(find_shape(slide, 138), f"{cur['ebc']:.0f} EBC")
    set_text(find_shape(slide, 140), f"{cur['sav']:.0f} SAV")
    set_delta(find_shape(slide, 146), f"N-1 : {n1['total']:.0f}", e["deplacement_total_n1"])
    set_delta(find_shape(slide, 142), f"N-1 : {n1['ebc']:.0f}", e["deplacement_ebc_n1"])
    set_delta(find_shape(slide, 134), f"N-1 : {n1['sav']:.0f}", e["deplacement_sav_n1"])


def _pct_abs(value) -> str:
    return f"{fr(value, 0)}%" if value is not None else "n/a"


def _fill_productivite(slide, result: dict):
    p, e = result["data"]["cur"]["productivite"], result["evolutions"]
    trav, inter, pct = p["heures_travaillees"], p["heures_intervention"], p["pct"]
    table = find_shape(slide, 9).table
    for col, equipe in ((1, "ebc"), (2, "sav"), (3, "total")):
        set_cell(table.cell(1, col), fr(trav[equipe], 0))
        set_cell(table.cell(2, col), _heures(inter[equipe]))
        pct_cell = table.cell(3, col)
        set_cell(pct_cell, _pct_abs(pct[equipe]))
        if pct[equipe] is not None:
            pct_cell.text_frame.paragraphs[0].runs[0].font.color.rgb = GREEN if pct[equipe] >= PRODUCTIVITE_SEUIL else RED

    set_text(find_shape(slide, 42), _pct_abs(pct["total"]))
    set_text(find_shape(slide, 52), f"{_pct_abs(pct['ebc'])} EBC")
    set_text(find_shape(slide, 56), f"{_pct_abs(pct['sav'])} SAV")

    m1 = (result.get("mois1_reference") or {}).get("productivite", {})
    _set_productivite_total_mois1(find_shape(slide, 41), m1.get("total"), e.get("productivite_total_m1"))
    for shape_id, champ, donor_id in ((58, "ebc", 55), (47, "sav", 50)):
        shape = find_shape(slide, shape_id)
        _ensure_second_paragraph(shape, find_shape(slide, donor_id))
        if m1.get(champ) is not None:
            set_delta(shape, f"Mois-1 : {_pct_abs(m1[champ])}", e.get(f"productivite_{champ}_m1"))
        else:
            set_delta_na(shape, "Mois-1")


def _fill_nb_interventions(slide, result: dict):
    d, e = result["data"]["cur"]["nb_interventions"], result["evolutions"]
    m1 = (result.get("mois1_reference") or {}).get("nb_interventions", {})

    set_text(find_shape(slide, 34), f"{d['total']} inter")
    set_text(find_shape(slide, 51), f"{d['ebc']} EBC")
    set_text(find_shape(slide, 53), f"{d['sav']} SAV")

    _set_interventions_total_mois1(find_shape(slide, 28), m1.get("total"), e.get("nb_interventions_total_m1"))
    for shape_id, champ in ((55, "ebc"), (50, "sav")):
        shape = find_shape(slide, shape_id)
        if m1.get(champ) is not None:
            set_delta(shape, f"Mois-1 : {m1[champ]}", e[f"nb_interventions_{champ}_m1"])
        else:
            set_delta_na(shape, "Mois-1")


def _set_simple_delta(shape, pct):
    """Boîte 'vs Mois-1' vide (aucun run à hériter, contrairement aux N-1/Mois-1 à 2 lignes des autres
    blocs) : juste la flèche + le %, ou 'n/a' si pas de référence Mois-1 fournie."""
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.font.size = Pt(12)
    if pct is None:
        run.text, run.font.color.rgb = "n/a", GREY
    else:
        run.text, run.font.color.rgb = f"{arrow(pct)} {fr_pct(pct)}", delta_color(pct)


def _fill_valorisation_stock(slide, result: dict):
    v, e = result["data"]["cur"]["valorisation_stock"], result["evolutions"]
    by_famille = {row["famille"]: row for row in v["rows"]}
    table = find_shape(slide, 27).table
    for i, (label, _) in enumerate(STOCK_FAMILLES, start=1):
        row = by_famille[label]
        set_cell(table.cell(i, 1), fr_k(row["depot"], 1))
        set_cell(table.cell(i, 2), fr_k(row["showroom"], 1))
        set_cell(table.cell(i, 3), fr_k(row["total"], 1))

    total_row = len(STOCK_FAMILLES) + 1
    set_cell(table.cell(total_row, 1), fr_k(v["total_depot"], 1))
    set_cell(table.cell(total_row, 2), fr_k(v["total_showroom"], 1))
    set_cell(table.cell(total_row, 3), fr_k(v["total"], 1))

    _set_simple_delta(find_shape(slide, 10), e.get("valorisation_stock_total_m1"))


def _fill_articles_epuisement(slide, result: dict):
    """Shape 20 est une zone de fond vide (comme la 9 de la Valorisation du stock, jamais remplie) : la
    valeur totale est en fait dans la 29 ('213 u'). Les 2 ratios (32, 35) ont déjà un libellé + une valeur
    de référence à écraser : 2e paragraphe pour 32, dernier run (après un retour à la ligne) pour 35."""
    a, e = result["data"]["cur"]["articles_epuisement"], result["evolutions"]

    set_text(find_shape(slide, 29), f"{a['total']} u")
    set_paragraph(find_shape(slide, 32).text_frame.paragraphs[1], _pct_abs(a["pct_ventes"]))
    find_shape(slide, 35).text_frame.paragraphs[0].runs[-1].text = _pct_abs(a["pct_dispo"])

    _set_simple_delta(find_shape(slide, 15), e.get("articles_epuisement_total_m1"))


def generate_sav_pptx(result: dict, template_path: str, output_path: str) -> str:
    prs = Presentation(template_path)
    slide1 = prs.slides[0]
    set_subtitle(find_shape(slide1, 5), result["year"], result["month"])
    _fill_devis(slide1, result)
    _fill_docs_nuls(slide1, result)
    _fill_main_oeuvre(slide1, result)
    _fill_deplacement(slide1, result)
    _fill_productivite(slide1, result)
    _fill_nb_interventions(slide1, result)

    # Le sous-titre de la diapo Achat/Appro (shape 5) n'a qu'un seul run (pas le schéma standard à 3
    # runs de set_subtitle) : le reste de cette diapo n'étant pas encore défini avec Antoine, on laisse
    # son texte tel quel.
    slide2 = prs.slides[1]
    _fill_valorisation_stock(slide2, result)
    _fill_articles_epuisement(slide2, result)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path
