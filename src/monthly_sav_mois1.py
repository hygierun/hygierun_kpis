"""Lecture du PowerPoint SAV/Achat du mois précédent pour le Mois-1 des blocs pas recalculables depuis
l'historique : Nombre d'interventions (diapo SAV), Valorisation du stock et Articles à épuisement (diapo
Achat/Appro). Comme pour Livraison & Compta, ces KPI viennent de fiches papier ou d'une photo du jour
(stock) — leur évolution Mois-1 se récupère en relisant le deck généré le mois précédent par cette appli.
"""

import re
from typing import Dict, Optional

from pptx import Presentation

from .pptx_helpers import find_shape

NB_INTERVENTIONS_SHAPES = {"total": 34, "ebc": 51, "sav": 53}
PRODUCTIVITE_SHAPES = {"total": 42, "ebc": 52, "sav": 56}
VALORISATION_STOCK_TOTAL_SHAPE = 34  # diapo Achat/Appro, ex. "749 k€"
ARTICLES_EPUISEMENT_TOTAL_SHAPE = 29  # diapo Achat/Appro, ex. "213 u"


def _parse_int(text: str) -> Optional[int]:
    match = re.search(r"(\d+)", text)
    return int(match.group(1)) if match else None


def _parse_pct(text: str) -> Optional[float]:
    match = re.search(r"(\d+(?:,\d+)?)\s*%", text)
    return float(match.group(1).replace(",", ".")) if match else None


def _parse_k_euros(text: str) -> Optional[float]:
    match = re.search(r"(\d+(?:,\d+)?)\s*k€", text)
    return float(match.group(1).replace(",", ".")) * 1000 if match else None


def read_sav_mois1_reference(pptx_path: str) -> Dict[str, dict]:
    """Lit le deck SAV/Achat autonome (2 diapos : SAV puis Achat/Appro)."""
    prs = Presentation(pptx_path)
    return _read_from_slides(prs.slides[0], prs.slides[1])


def _read_from_slides(slide_sav, slide_achat) -> Dict[str, dict]:
    """Cœur de la lecture, indépendant de la position des diapos dans le fichier."""
    nb_interventions = {
        champ: _parse_int(find_shape(slide_sav, shape_id).text_frame.text)
        for champ, shape_id in NB_INTERVENTIONS_SHAPES.items()
    }
    productivite = {
        champ: _parse_pct(find_shape(slide_sav, shape_id).text_frame.text)
        for champ, shape_id in PRODUCTIVITE_SHAPES.items()
    }
    valorisation_stock = {"total": _parse_k_euros(find_shape(slide_achat, VALORISATION_STOCK_TOTAL_SHAPE).text_frame.text)}
    articles_epuisement = {"total": _parse_int(find_shape(slide_achat, ARTICLES_EPUISEMENT_TOTAL_SHAPE).text_frame.text)}

    return {"nb_interventions": nb_interventions, "productivite": productivite,
            "valorisation_stock": valorisation_stock, "articles_epuisement": articles_epuisement}
